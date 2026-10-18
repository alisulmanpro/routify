from __future__ import annotations

import asyncio
import csv
import io
import zipfile
from datetime import date, datetime
from typing import Any, Optional

from ..errors import unprocessable
from ..uploads import UploadedFile, detect_document_type
from .pdf_import import extract_addresses_from_lines, extract_deliveries_from_text, extract_pdf_text, flag_duplicates
from .subscription_service import assert_can_create_deliveries


# A 10 MB upload may not unpack into gigabytes (zip bomb): real delivery sheets are a tiny fraction of these limits.
MAX_UNZIPPED_BYTES = 100 * 1024 * 1024
MAX_ZIP_ENTRIES = 2000
MAX_COMPRESSION_RATIO = 200


def _assert_safe_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
    except (zipfile.BadZipFile, ValueError):
        raise unprocessable("This file could not be read. Check that it is a valid Office document.", "INVALID_FILE") from None
    total = sum(e.file_size for e in entries)
    packed = sum(e.compress_size for e in entries) or 1
    if len(entries) > MAX_ZIP_ENTRIES or total > MAX_UNZIPPED_BYTES or total / packed > MAX_COMPRESSION_RATIO:
        raise unprocessable("This file is too large to import once unpacked. Split it into smaller files.", "FILE_TOO_COMPLEX")


def _extract_from_word(data: bytes, ext: str) -> str:
    if ext == "doc":
        raise unprocessable("The older .doc format is not supported. Please open the file in Word and save it as a .docx document.", "UNSUPPORTED_DOC")
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(io.BytesIO(data))
        parts: list[str] = []

        def walk(container) -> None:
            # Body order matters: paragraphs and tables are read as they appear, like mammoth's raw text.
            for child in container.iter_inner_content():
                if isinstance(child, Paragraph):
                    parts.append(child.text)
                elif isinstance(child, Table):
                    for row in child.rows:
                        for cell in row.cells:
                            walk(cell)

        walk(document)
        return "\n\n".join(parts) + ("\n\n" if parts else "")
    except Exception:  # noqa: BLE001
        raise unprocessable("This Word document could not be read. Check that it is a valid .docx file.", "INVALID_DOCX") from None


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _lines(rows) -> list[str]:
    out = []
    for row in rows:
        line = "  ".join(c for c in (_cell(v) for v in row) if c)
        if line:
            out.append(line)
    return out


def _extract_from_spreadsheet(data: bytes, ext: str) -> str:
    try:
        lines: list[str] = []
        if ext == "csv":
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = data.decode("latin-1")
            lines = _lines(csv.reader(io.StringIO(text)))
        elif ext == "xls":
            import xlrd

            book = xlrd.open_workbook(file_contents=data)
            for sheet in book.sheets():
                lines += _lines(sheet.row_values(r) for r in range(sheet.nrows))
        else:
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            for sheet in book.worksheets:
                lines += _lines(sheet.iter_rows(values_only=True))
            book.close()
        return "\n".join(lines)
    except Exception:  # noqa: BLE001
        raise unprocessable("This spreadsheet could not be read. Check that it is a valid Excel or CSV file.", "INVALID_SPREADSHEET") from None


def _extract_text(data: bytes, ext: str) -> str:
    if ext in ("docx", "xlsx"):
        _assert_safe_zip(data)
    if ext == "pdf":
        return extract_pdf_text(data)[0]
    if ext in ("docx", "doc"):
        return _extract_from_word(data, ext)
    if ext in ("xlsx", "xls", "csv"):
        return _extract_from_spreadsheet(data, ext)
    if ext == "txt":
        try:
            return data.decode("utf-8-sig")
        except UnicodeDecodeError:
            return data.decode("latin-1")
    raise unprocessable("Unsupported file type", "UNSUPPORTED_FILE")


async def build_draft_from_file(actor: dict, file: Optional[UploadedFile]) -> dict:
    ext = detect_document_type(file.data if file else None, file.filename if file else None)
    await assert_can_create_deliveries(actor, 1)

    text = await asyncio.to_thread(_extract_text, file.data, ext)
    if not text.strip():
        raise unprocessable("No text found in this file. Try a different file or add deliveries manually.", "FILE_NO_TEXT")
    # A courier manifest (parcel IDs, names, phones) first; otherwise a plain list with one address per line.
    items = extract_deliveries_from_text(text) or extract_addresses_from_lines(text)
    if not items:
        raise unprocessable("No deliveries were recognised in this file. Add them manually instead.", "FILE_NO_DELIVERIES")
    await flag_duplicates(actor, items)
    return {"fileType": ext, "items": items}
